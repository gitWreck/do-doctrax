import os
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import datetime

DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "doctrax.db")

def export_db_to_excel(output_path=None):
    if not output_path:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), f"DOCTRAX_EXPORT_{timestamp}.xlsx")

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="145335", end_color="145335", fill_type="solid") # CFNR Green
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    def style_header(ws, col_count):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. INCOMING-OUTGOING DOCS
    ws1 = wb.create_sheet(title="INCOMING-OUTGOING DOCS")
    headers1 = ["ID", "Date", "Doctrax", "From", "Document Type", "Description", "Status", "Received by", "Remarks"]
    ws1.append(headers1)
    style_header(ws1, len(headers1))
    cur.execute("SELECT id, date, doctrax_no, from_unit, doc_type, description, status, received_by, remarks FROM incoming_outgoing_docs ORDER BY id DESC")
    for row in cur.fetchall():
        ws1.append(list(row))

    # 2. EMAIL RECEIVED
    ws2 = wb.create_sheet(title="EMAIL RECEIVED")
    headers2 = ["ID", "Date", "From", "Doctrax", "Description", "Deadline/Link", "Remarks"]
    ws2.append(headers2)
    style_header(ws2, len(headers2))
    cur.execute("SELECT id, date, sender, doctrax_no, description, deadline_link, remarks FROM email_logs ORDER BY id DESC")
    for row in cur.fetchall():
        ws2.append(list(row))

    # 3. MEMO, AO RECEIVED IN EMAIL
    ws3 = wb.create_sheet(title="MEMO, AO RECEIVED IN EMAIL")
    headers3 = ["ID", "Date", "From", "Document Type", "Memo No.", "Description / Subject", "Remarks"]
    ws3.append(headers3)
    style_header(ws3, len(headers3))
    cur.execute("SELECT id, date, sender, doc_type, memo_no, subject, remarks FROM memo_ao_external ORDER BY id DESC")
    for row in cur.fetchall():
        ws3.append(list(row))

    # 4. APPOINTMENT, NOSA, ETC.
    ws4 = wb.create_sheet(title="APPOINTMENT, NOSA, ETC.")
    headers4 = ["ID", "Date", "Doctrack", "From", "To", "Unit", "Document Type", "Description", "Received by", "Received by UNITS", "Remarks"]
    ws4.append(headers4)
    style_header(ws4, len(headers4))
    cur.execute("SELECT id, date, doctrack_no, from_office, to_person, unit, doc_type, description, received_by, received_by_units, remarks FROM personnel_appointments ORDER BY id DESC")
    for row in cur.fetchall():
        ws4.append(list(row))

    # 5. CFNR MEMO
    ws5 = wb.create_sheet(title="CFNR MEMO")
    headers5 = ["ID", "Date", "CFNR Memo No.", "To", "UNIT FOR OIC", "Document Description", "Remarks"]
    ws5.append(headers5)
    style_header(ws5, len(headers5))
    cur.execute("SELECT id, date, memo_no, recipient, unit_for_oic, description, remarks FROM cfnr_memos ORDER BY id DESC")
    for row in cur.fetchall():
        ws5.append(list(row))

    # 6. CFNR ADMIN ORDER
    ws6 = wb.create_sheet(title="CFNR ADMIN ORDER")
    headers6 = ["ID", "Date", "CFNR AO No.", "To", "Document Description", "Remarks"]
    ws6.append(headers6)
    style_header(ws6, len(headers6))
    cur.execute("SELECT id, date, ao_no, recipient, description, remarks FROM cfnr_admin_orders ORDER BY id DESC")
    for row in cur.fetchall():
        ws6.append(list(row))

    # 7. OUT PPMP
    ws7 = wb.create_sheet(title="OUT PPMP")
    headers7 = ["ID", "Date", "Document Type/Description", "Doctrack No.", "Amount", "Remarks"]
    ws7.append(headers7)
    style_header(ws7, len(headers7))
    cur.execute("SELECT id, date, description, doctrack_no, amount, remarks FROM ppmp_procurement ORDER BY id DESC")
    for row in cur.fetchall():
        ws7.append(list(row))

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col[:50])
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(output_path)
    conn.close()
    print(f"Exported successfully to: {output_path}")
    return output_path

if __name__ == "__main__":
    export_db_to_excel()
