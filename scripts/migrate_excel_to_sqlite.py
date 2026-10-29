import os
import sys
import glob
import sqlite3
import openpyxl
import datetime

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, "doctrax.db")

def format_date(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if len(s) >= 10 and s[4] == '-' and s[7] == '-':
        return s[:10]
    return s

def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()

def init_db(conn):
    cur = conn.cursor()
    
    # 1. Incoming & Outgoing Docs
    cur.execute("""
    CREATE TABLE IF NOT EXISTS incoming_outgoing_docs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        doctrax_no TEXT,
        from_unit TEXT,
        doc_type TEXT,
        description TEXT,
        status TEXT,
        received_by TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_in_out_date ON incoming_outgoing_docs(date);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_in_out_doctrax ON incoming_outgoing_docs(doctrax_no);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_in_out_unit ON incoming_outgoing_docs(from_unit);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_in_out_status ON incoming_outgoing_docs(status);")

    # 2. Email Logs
    cur.execute("""
    CREATE TABLE IF NOT EXISTS email_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        sender TEXT,
        doctrax_no TEXT,
        description TEXT,
        deadline_link TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_email_date ON email_logs(date);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_email_sender ON email_logs(sender);")

    # 3. External Memos & Admin Orders
    cur.execute("""
    CREATE TABLE IF NOT EXISTS memo_ao_external (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        sender TEXT,
        doc_type TEXT,
        memo_no TEXT,
        subject TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_memo_ext_date ON memo_ao_external(date);")

    # 4. Personnel & Appointments
    cur.execute("""
    CREATE TABLE IF NOT EXISTS personnel_appointments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        doctrack_no TEXT,
        from_office TEXT,
        to_person TEXT,
        unit TEXT,
        doc_type TEXT,
        description TEXT,
        received_by TEXT,
        received_by_units TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_appt_to ON personnel_appointments(to_person);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_appt_unit ON personnel_appointments(unit);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_appt_type ON personnel_appointments(doc_type);")

    # 5. CFNR Internal Memos
    cur.execute("""
    CREATE TABLE IF NOT EXISTS cfnr_memos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        memo_no TEXT,
        recipient TEXT,
        unit_for_oic TEXT,
        description TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 6. CFNR Internal Admin Orders
    cur.execute("""
    CREATE TABLE IF NOT EXISTS cfnr_admin_orders (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        ao_no TEXT,
        recipient TEXT,
        description TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 7. Outgoing PPMP
    cur.execute("""
    CREATE TABLE IF NOT EXISTS ppmp_procurement (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        date TEXT,
        description TEXT,
        doctrack_no TEXT,
        amount TEXT,
        remarks TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 8. Controlled Vocabulary Lookups
    cur.execute("""
    CREATE TABLE IF NOT EXISTS lookup_values (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category TEXT,
        value TEXT,
        sort_order INTEGER DEFAULT 0,
        UNIQUE(category, value) ON CONFLICT IGNORE
    );
    """)
    conn.commit()

def iterate_valid_rows(ws, start_row=1, max_empty=60):
    empty_count = 0
    for idx, r in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
        is_empty = not any(c is not None and str(c).strip() != '' for c in r)
        if is_empty:
            empty_count += 1
            if empty_count >= max_empty:
                break
            continue
        empty_count = 0
        yield r

def migrate_workbook(excel_path, cur):
    print(f"\n==================================================")
    print(f"Loading Workbook: {os.path.basename(excel_path)}")
    print(f"==================================================")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

    # 1. Lookup values (Sheet1)
    if "Sheet1" in wb.sheetnames:
        ws = wb["Sheet1"]
        doc_types = []
        appt_types = []
        memo_types = []
        statuses = []
        units = []
        
        for row in iterate_valid_rows(ws, start_row=1):
            if row[0]: doc_types.append(clean_str(row[0]))
            if len(row) > 1 and row[1]:
                val = clean_str(row[1])
                if val in ["DO", "OCS", "OCREL", "LIB", "SFFG", "FBS", "FPPS", "IRNR", "FDC", "IAF", "MCME", "TREES", "INREM", "N/A"]:
                    units.append(val)
                else:
                    appt_types.append(val)
            if len(row) > 2 and row[2]: memo_types.append(clean_str(row[2]))
            if len(row) > 3 and row[3]: statuses.append(clean_str(row[3]))

        for i, val in enumerate(doc_types):
            cur.execute("INSERT OR IGNORE INTO lookup_values (category, value, sort_order) VALUES ('doc_type', ?, ?)", (val, i))
        for i, val in enumerate(appt_types):
            cur.execute("INSERT OR IGNORE INTO lookup_values (category, value, sort_order) VALUES ('appointment_type', ?, ?)", (val, i))
        for i, val in enumerate(units):
            cur.execute("INSERT OR IGNORE INTO lookup_values (category, value, sort_order) VALUES ('unit', ?, ?)", (val, i))
        for i, val in enumerate(statuses):
            cur.execute("INSERT OR IGNORE INTO lookup_values (category, value, sort_order) VALUES ('status', ?, ?)", (val, i))
        for i, val in enumerate(memo_types):
            cur.execute("INSERT OR IGNORE INTO lookup_values (category, value, sort_order) VALUES ('memo_type', ?, ?)", (val, i))

    for sname in wb.sheetnames:
        sname_upper = sname.upper().strip()

        # 2. INCOMING-OUTGOING DOCS
        if sname_upper.startswith("INCOMING-OUTGOING"):
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=3):
                date_val = format_date(r[1]) if len(r) > 1 else ""
                doctrax_val = clean_str(r[2]) if len(r) > 2 else ""
                from_val = clean_str(r[3]) if len(r) > 3 else ""
                type_val = clean_str(r[4]) if len(r) > 4 else ""
                desc_val = clean_str(r[5]) if len(r) > 5 else ""
                status_val = clean_str(r[6]) if len(r) > 6 else ""
                rec_val = clean_str(r[7]) if len(r) > 7 else ""
                rem_val = clean_str(r[8]) if len(r) > 8 else ""

                if date_val.lower() == 'date' or from_val.lower() == 'from' or (not date_val and not desc_val and not doctrax_val):
                    continue

                cur.execute("""
                INSERT INTO incoming_outgoing_docs (date, doctrax_no, from_unit, doc_type, description, status, received_by, remarks)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (date_val, doctrax_val, from_val, type_val, desc_val, status_val, rec_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 3. EMAIL RECEIVED
        elif sname_upper.startswith("EMAIL RECEIVED"):
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=3):
                date_val = format_date(r[1]) if len(r) > 1 else ""
                sender_val = clean_str(r[2]) if len(r) > 2 else ""
                doctrax_val = clean_str(r[3]) if len(r) > 3 else ""
                desc_val = clean_str(r[4]) if len(r) > 4 else ""
                deadline_val = clean_str(r[5]) if len(r) > 5 else ""
                rem_val = clean_str(r[6]) if len(r) > 6 else ""

                if date_val.lower() == 'date' or sender_val.lower() == 'from' or (not date_val and not desc_val and not sender_val):
                    continue

                cur.execute("""
                INSERT INTO email_logs (date, sender, doctrax_no, description, deadline_link, remarks)
                VALUES (?, ?, ?, ?, ?, ?)
                """, (date_val, sender_val, doctrax_val, desc_val, deadline_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 4. MEMO, AO RECEIVED IN EMAIL
        elif sname_upper.startswith("MEMO, AO"):
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=3):
                date_val = format_date(r[1]) if len(r) > 1 else ""
                sender_val = clean_str(r[2]) if len(r) > 2 else ""
                type_val = clean_str(r[3]) if len(r) > 3 else ""
                memo_no_val = clean_str(r[4]) if len(r) > 4 else ""
                subj_val = clean_str(r[5]) if len(r) > 5 else ""
                rem_val = clean_str(r[6]) if len(r) > 6 else ""

                if date_val.lower() == 'date' or sender_val.lower() == 'from' or (not date_val and not subj_val and not memo_no_val):
                    continue

                cur.execute("""
                INSERT INTO memo_ao_external (date, sender, doc_type, memo_no, subject, remarks)
                VALUES (?, ?, ?, ?, ?, ?)
                """, (date_val, sender_val, type_val, memo_no_val, subj_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 5. APPOINTMENT, NOSA, ETC.
        elif sname_upper.startswith("APPOINTMENT"):
            ws = wb[sname]
            count = 0
            # Inspect column layout in row 1-3
            has_doctrack_col = True
            for check_row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                check_vals = [clean_str(c).lower() for c in check_row if c is not None]
                if 'from' in check_vals and 'to' in check_vals:
                    row_clean = [clean_str(c).lower() for c in check_row]
                    if len(row_clean) > 1 and row_clean[1] == 'from':
                        has_doctrack_col = False
                    break

            for r in iterate_valid_rows(ws, start_row=2):
                date_val = format_date(r[0]) if len(r) > 0 else ""
                if date_val.lower() == 'date' or date_val.lower().startswith('appointment'):
                    continue

                if has_doctrack_col:
                    doctrack_val = clean_str(r[1]) if len(r) > 1 else ""
                    from_val = clean_str(r[2]) if len(r) > 2 else ""
                    to_val = clean_str(r[3]) if len(r) > 3 else ""
                    unit_val = clean_str(r[4]) if len(r) > 4 else ""
                    type_val = clean_str(r[5]) if len(r) > 5 else ""
                    desc_val = clean_str(r[6]) if len(r) > 6 else ""
                    rec_val = clean_str(r[7]) if len(r) > 7 else ""
                    rec_unit_val = clean_str(r[8]) if len(r) > 8 else ""
                    rem_val = clean_str(r[9]) if len(r) > 9 else ""
                else:
                    doctrack_val = ""
                    from_val = clean_str(r[1]) if len(r) > 1 else ""
                    to_val = clean_str(r[2]) if len(r) > 2 else ""
                    unit_val = clean_str(r[3]) if len(r) > 3 else ""
                    type_val = clean_str(r[4]) if len(r) > 4 else ""
                    desc_val = clean_str(r[5]) if len(r) > 5 else ""
                    rec_val = clean_str(r[6]) if len(r) > 6 else ""
                    rec_unit_val = ""
                    rem_val = clean_str(r[7]) if len(r) > 7 else ""

                if not date_val and not to_val and not desc_val:
                    continue

                cur.execute("""
                INSERT INTO personnel_appointments (date, doctrack_no, from_office, to_person, unit, doc_type, description, received_by, received_by_units, remarks)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (date_val, doctrack_val, from_val, to_val, unit_val, type_val, desc_val, rec_val, rec_unit_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 6. CFNR MEMO
        elif sname_upper.startswith("CFNR MEMO"):
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=2):
                date_val = format_date(r[0]) if len(r) > 0 else ""
                memo_no_val = clean_str(r[1]) if len(r) > 1 else ""
                to_val = clean_str(r[2]) if len(r) > 2 else ""
                unit_val = clean_str(r[3]) if len(r) > 3 else ""
                desc_val = clean_str(r[4]) if len(r) > 4 else ""
                rem_val = clean_str(r[5]) if len(r) > 5 else ""

                if date_val.lower() == 'date' or memo_no_val.lower().startswith('cfnr memo') or (not date_val and not to_val and not desc_val):
                    continue

                cur.execute("""
                INSERT INTO cfnr_memos (date, memo_no, recipient, unit_for_oic, description, remarks)
                VALUES (?, ?, ?, ?, ?, ?)
                """, (date_val, memo_no_val, to_val, unit_val, desc_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 7. CFNR ADMIN ORDER
        elif sname_upper.startswith("CFNR ADMIN ORDER"):
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=2):
                date_val = format_date(r[0]) if len(r) > 0 else ""
                ao_no_val = clean_str(r[1]) if len(r) > 1 else ""
                to_val = clean_str(r[2]) if len(r) > 2 else ""
                desc_val = clean_str(r[3]) if len(r) > 3 else ""
                rem_val = clean_str(r[4]) if len(r) > 4 else ""

                if date_val.lower() == 'date' or ao_no_val.lower().startswith('cfnr admin') or (not date_val and not to_val and not desc_val):
                    continue

                cur.execute("""
                INSERT INTO cfnr_admin_orders (date, ao_no, recipient, description, remarks)
                VALUES (?, ?, ?, ?, ?)
                """, (date_val, ao_no_val, to_val, desc_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 8. OUT PPMP
        elif "PPMP" in sname_upper:
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=2):
                date_val = format_date(r[0]) if len(r) > 0 else ""
                desc_val = clean_str(r[1]) if len(r) > 1 else ""
                doctrack_val = clean_str(r[2]) if len(r) > 2 else ""
                rem_val = clean_str(r[3]) if len(r) > 3 else ""

                if date_val.lower() == 'date' or (not date_val and not desc_val):
                    continue

                cur.execute("""
                INSERT INTO ppmp_procurement (date, description, doctrack_no, amount, remarks)
                VALUES (?, ?, ?, ?, ?)
                """, (date_val, desc_val, doctrack_val, "", rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

        # 9. GFOREST
        elif sname_upper == "GFOREST":
            ws = wb[sname]
            count = 0
            for r in iterate_valid_rows(ws, start_row=2):
                date_val = format_date(r[0]) if len(r) > 0 else ""
                doctrax_val = clean_str(r[1]) if len(r) > 1 else ""
                type_val = clean_str(r[2]) if len(r) > 2 else ""
                desc_val = clean_str(r[3]) if len(r) > 3 else ""
                status_val = clean_str(r[4]) if len(r) > 4 else ""
                rec_val = clean_str(r[5]) if len(r) > 5 else ""
                rem_val = clean_str(r[6]) if len(r) > 6 else ""

                if date_val.lower() == 'date' or (not date_val and not desc_val):
                    continue

                cur.execute("""
                INSERT INTO incoming_outgoing_docs (date, doctrax_no, from_unit, doc_type, description, status, received_by, remarks)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (date_val, doctrax_val, "GFOREST", type_val, desc_val, status_val, rec_val, rem_val))
                count += 1
            print(f"  [+] {sname}: {count} records imported")

    wb.close()

def migrate():
    # Find all available Excel workbooks
    if len(sys.argv) > 1:
        workbooks = [os.path.abspath(a) for a in sys.argv[1:]]
    else:
        files = glob.glob(os.path.join(BASE_DIR, "COMMUNICATIONS & DOCUMENTS RECEIVED *.xlsx"))
        # Sort so 2025 comes before 2026
        workbooks = sorted(files)

    if not workbooks:
        print("No Excel files found to migrate.")
        return

    print(f"Opening database: {DB_PATH}")
    if os.path.exists(DB_PATH):
        try:
            os.remove(DB_PATH)
            print("Removed existing database to perform clean seed of all years.")
        except Exception as e:
            print(f"Notice: {e}")

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL;")
    init_db(conn)
    cur = conn.cursor()

    for wb_path in workbooks:
        migrate_workbook(wb_path, cur)
        conn.commit()

    # Total counts summary
    print(f"\n==================================================")
    print("MIGRATION COMPLETE: SUMMARY OF RECORDS")
    print(f"==================================================")
    tables = [
        ("incoming_outgoing_docs", "Incoming & Outgoing Docs"),
        ("email_logs", "Email Received"),
        ("memo_ao_external", "Memos & AO External"),
        ("personnel_appointments", "Personnel & Appointments"),
        ("cfnr_memos", "CFNR Memos"),
        ("cfnr_admin_orders", "CFNR Admin Orders"),
        ("ppmp_procurement", "PPMP Procurement")
    ]
    total_all = 0
    for tbl, lbl in tables:
        cur.execute(f"SELECT COUNT(*) FROM {tbl}")
        cnt = cur.fetchone()[0]
        total_all += cnt
        print(f"  {lbl.ljust(26)}: {cnt:>6} records")
    print(f"--------------------------------------------------")
    print(f"  {'TOTAL ALL RECORDS'.ljust(26)}: {total_all:>6} records")
    print(f"==================================================\n")

    conn.commit()
    conn.close()

if __name__ == "__main__":
    migrate()
